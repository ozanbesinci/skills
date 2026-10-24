"""Mail eklerinden metin çıkarır (PDF, XLSX, DOCX) → index\\ekler.sqlite. Arşive yalnız okur.

    py -3 -E ek_cikar.py           # bütün aylar, en yeniden geriye; artımlı, kaldığı yerden
    py -3 -E ek_cikar.py --ay 2026-09
    py -3 -E ek_cikar.py --durum

SİSTEM Python'uyla ve -I OLMADAN koşar: PyMuPDF, openpyxl, python-docx sistem Python'unda
kurulu (openpyxl ve python-docx kullanıcı paketlerinde; -I onları gizler), mail-ara'nın sanal
ortamında yok. İndeksleme ayrı adımdır: ek_indeksle.py (sanal ortam).
Metin katmanı olmayan (taranmış) PDF "metin_yok" olarak işaretlenir; OCR bu sürümde yok.
Aynı içerik (SHA-1) bir kez saklanır; hangi maillerde geçtiği `ek` tablosundadır.
"""
import aktarim
import argparse
import hashlib
import os
import sqlite3
import sys
import time
import zipfile
from pathlib import Path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ortak

KOK = Path(os.environ.get("MAIL_INDEKS_KOK", r"D:\AI-BACKUP\mail-indeks"))
ARSIV = Path(os.environ.get("MAIL_ARSIV_KOK", str(Path.home() / "Drive'ım" / "maillerim")))
DB = KOK / "index" / "ekler.sqlite"
KILIT = KOK / "index" / "indeksle.kilit"
AYLAR = ["ocak", "subat", "mart", "nisan", "mayis", "haziran",
         "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]
TURLER = {".pdf", ".xlsx", ".xlsm", ".docx"}
MAKS_KARAKTER, MAKS_SAYFA, MAKS_SATIR = 60000, 40, 3000

SEMA = """
CREATE TABLE IF NOT EXISTS ek(dosya TEXT PRIMARY KEY, mail_id TEXT, ay TEXT, ad TEXT,
    uzanti TEXT, boyut INTEGER, mtime REAL, ozet TEXT, durum TEXT);
CREATE INDEX IF NOT EXISTS ek_mail ON ek(mail_id);
CREATE INDEX IF NOT EXISTS ek_ozet ON ek(ozet);
CREATE TABLE IF NOT EXISTS icerik(ozet TEXT PRIMARY KEY, metin TEXT, karakter INTEGER,
    sayfa INTEGER, durum TEXT);
CREATE TABLE IF NOT EXISTS ay_durum(ay TEXT PRIMARY KEY, dosya INTEGER, zaman TEXT);
"""


def ay_klasoru(ay):
    return aktarim.ay_klasoru(ARSIV, ay)


def arsiv_aylari():
    return list(reversed(aktarim.arsiv_aylari(ARSIV)))


def pdf_metin(yol):
    import fitz
    with fitz.open(yol) as doc:
        if doc.needs_pass:
            return "", 0, "sifreli"
        parca, n, dolu = [], 0, 0
        for i, sayfa in enumerate(doc):
            if i >= MAKS_SAYFA or n >= MAKS_KARAKTER:
                break
            t = sayfa.get_text("text")
            # Sayfa numarası etiketleri taranmış PDF'yi metinli gösteremez.
            dolu += sum(1 for c in t if c.isalnum())
            parca.append(f"[s.{i + 1}]\n{t}")
            n += len(t)
        metin = "\n".join(parca)[:MAKS_KARAKTER]
        return metin, doc.page_count, ("ok" if dolu > 40 else "metin_yok")


def xlsx_metin(yol):
    import openpyxl
    wb = openpyxl.load_workbook(yol, read_only=True, data_only=True)
    try:
        parca, n = [], 0
        for ws in wb.worksheets:
            parca.append(f"[sayfa: {ws.title}]")
            for i, satir in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAKS_SATIR or n >= MAKS_KARAKTER:
                    break
                hucre = [str(v) for v in satir if v not in (None, "")]
                if hucre:
                    s = " | ".join(hucre)
                    parca.append(s)
                    n += len(s)
        metin = "\n".join(parca)[:MAKS_KARAKTER]
        return metin, len(wb.worksheets), ("ok" if n else "metin_yok")
    finally:
        wb.close()


def docx_metin(yol):
    import docx
    d = docx.Document(yol)
    parca = [p.text for p in d.paragraphs if p.text.strip()]
    for tablo in d.tables:
        for satir in tablo.rows:
            parca.append(" | ".join(c.text.strip() for c in satir.cells if c.text.strip()))
    metin = "\n".join(parca)[:MAKS_KARAKTER]
    return metin, 0, ("ok" if any(c.isalnum() for c in metin) else "metin_yok")


def cikar(yol):
    uz = yol.suffix.lower()
    try:
        if uz == ".pdf":
            return pdf_metin(str(yol))
        if uz in (".xlsx", ".xlsm"):
            return xlsx_metin(str(yol))
        if uz == ".docx":
            return docx_metin(str(yol))
    except zipfile.BadZipFile:
        return "", 0, "bozuk"
    except Exception as e:  # bozuk/garip dosya bütün koşuyu durdurmasın
        return f"{type(e).__name__}: {str(e)[:200]}", 0, "hata"
    return "", 0, "atlandi"


def ay_isle(db, ay):
    aktarim.stage(KOK, ARSIV, ay, "ek_cikar")
    generation = aktarim.ensure_generation(KOK, "ek_cikar")
    ekler = ay_klasoru(ay) / "ekler"
    live_paths = set()
    # Hata almış kayıtlar her koşuda yeniden denenir (ör. eksik kütüphane düzeltildikten sonra).
    kayitli = {r[0]: (r[1], r[2]) for r in db.execute(
        "SELECT dosya, boyut, mtime FROM ek WHERE ay=? AND durum != 'hata'", (ay,))}
    bilinen = {r[0] for r in db.execute("SELECT ozet FROM icerik WHERE durum != 'hata'")}
    yeni = cikan = 0
    t0 = time.time()
    for mail_dir in sorted(p for p in ekler.iterdir() if p.is_dir()):
        mail_id = mail_dir.name.split("_", 1)[0]
        for yol in mail_dir.rglob("*"):
            if not yol.is_file() or yol.suffix.lower() not in TURLER:
                continue
            st = yol.stat()
            anahtar = str(yol)
            live_paths.add(anahtar)
            ozet = hashlib.sha1(yol.read_bytes()).hexdigest()
            previous = db.execute("SELECT ozet, durum FROM ek WHERE dosya=?", (anahtar,)).fetchone()
            if previous and previous[0] == ozet and previous[1] != "hata":
                continue
            if ozet not in bilinen:
                metin, sayfa, durum = cikar(yol)
                db.execute("INSERT OR REPLACE INTO icerik VALUES (?,?,?,?,?)",
                           (ozet, metin if durum in ("ok", "metin_yok") else "", len(metin), sayfa, durum))
                bilinen.add(ozet)
                cikan += 1
            else:
                durum = db.execute("SELECT durum FROM icerik WHERE ozet=?", (ozet,)).fetchone()[0]
            db.execute("INSERT OR REPLACE INTO ek VALUES (?,?,?,?,?,?,?,?,?)",
                       (anahtar, mail_id, ay, yol.name, yol.suffix.lower(), st.st_size, st.st_mtime, ozet, durum))
            yeni += 1
            if yeni % 500 == 0:
                db.commit()
                print(f"  {ay}: {yeni} ek ({cikan} yeni içerik), {time.time() - t0:.0f} sn", flush=True)
    stale = {r[0] for r in db.execute("SELECT dosya FROM ek WHERE ay=?", (ay,))} - live_paths
    db.executemany("DELETE FROM ek WHERE dosya=?", [(p,) for p in stale])
    db.execute("DELETE FROM icerik WHERE ozet NOT IN (SELECT DISTINCT ozet FROM ek)")
    db.execute("INSERT OR REPLACE INTO ay_durum VALUES (?,?,datetime('now','localtime'))",
               (ay, db.execute("SELECT COUNT(*) FROM ek WHERE ay=?", (ay,)).fetchone()[0]))
    db.commit()
    actual = {r[0] for r in db.execute("SELECT dosya FROM ek WHERE ay=?", (ay,))}
    if actual != live_paths:
        raise ValueError("Ek kaynak ilişkileri uzlaştırılamadı")
    statuses = dict(db.execute("SELECT durum, COUNT(*) FROM ek WHERE ay=? GROUP BY durum", (ay,)))
    if statuses.get("hata"):
        raise ValueError("Ek çıkarma hataları var; teslim tamamlanmadı")
    aktarim.stage(KOK, ARSIV, ay, "ek_cikar", True, dict(generation=generation, durumlar=statuses))
    print(f"{ay}: {yeni} ek işlendi, {cikan} yeni içerik, {time.time() - t0:.0f} sn", flush=True)


def durum(db):
    for ay, n, z in db.execute("SELECT * FROM ay_durum ORDER BY ay"):
        print(f"{ay}  {n:>6} ek  {z}")
    for d, n in db.execute("SELECT durum, COUNT(*) FROM icerik GROUP BY durum ORDER BY 2 DESC"):
        print(f"  içerik {d:10} {n}")


def main():
    ortak.utf8_cikti()
    ap = argparse.ArgumentParser()
    ap.add_argument("--ay", action="append", default=[])
    ap.add_argument("--durum", action="store_true")
    a = ap.parse_args()
    DB.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB, timeout=60)
    db.executescript(SEMA)
    if a.durum:
        durum(db)
        return
    ortak.kilit_al(KILIT)
    if os.name == "nt":
        import ctypes
        ctypes.windll.kernel32.SetThreadExecutionState(0x80000000 | 0x00000001)
    try:
        for ay in (a.ay or arsiv_aylari()):
            ay_isle(db, ay)
    finally:
        db.close()
        KILIT.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
