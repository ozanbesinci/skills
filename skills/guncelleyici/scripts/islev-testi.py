"""program-guncelle — islev testi.

kurulum-testi.py tarafindan, kurulan kutuphanelerle ayni Python'da alt surec olarak calistirilir.
Her kutuphaneyi gercek bir isle dener ve her sonucu tek satirda "@@T {json}" olarak basar.
Bir testin cokmesi digerlerini durdurmaz. Yalniz verilen calisma klasorune yazar.

Kullanim: python islev-testi.py <calisma-klasoru>
"""
import json
import os
import sys
import time
import traceback

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # noqa: BLE001
    pass

CALISMA = sys.argv[1]
os.makedirs(CALISMA, exist_ok=True)
METIN = "PROGRAM GUNCELLE TESTI 12345"


def bildir(kalem, ad, durum, ayrinti):
    print("@@T " + json.dumps({"kalem": kalem, "ad": ad, "durum": durum, "ayrinti": ayrinti},
                              ensure_ascii=False), flush=True)


def hata_ozeti(exc):
    son = traceback.format_exception_only(type(exc), exc)[-1].strip()
    return son[-400:]


def test(kalem, ad):
    def sarici(fn):
        try:
            ayrinti = fn()
            bildir(kalem, ad, "GEÇTİ", ayrinti)
            return True
        except Exception as exc:  # noqa: BLE001 — her hata rapora gitmeli
            bildir(kalem, ad, "ENGEL", hata_ozeti(exc))
            return False
    return sarici


def surum(paket):
    from importlib.metadata import version
    try:
        return version(paket)
    except Exception:  # noqa: BLE001
        return "?"


# ------------------------------------------------------------------ kutuphane yukleme
K2 = "Kalem 2 - Kütüphane kurulumu"
K3 = "Kalem 3 - Excel otomasyonu"

yuklenen = {}
for modul, paket in [("openpyxl", "openpyxl"), ("pypdf", "pypdf"), ("pdfplumber", "pdfplumber"),
                     ("pypdfium2", "pypdfium2"), ("PIL.Image", "pillow"),
                     ("cryptography.hazmat.primitives.hashes", "cryptography"),
                     ("pythoncom", "pywin32"), ("win32com.client", "pywin32")]:
    @test(K2, f"{modul} yükleme")
    def _yukle(modul=modul, paket=paket):
        __import__(modul)
        yuklenen[modul] = True
        return f"{paket} {surum(paket)} yüklendi"


# ------------------------------------------------------------------ islev testleri
xlsx = os.path.join(CALISMA, "deneme.xlsx")
pdf = os.path.join(CALISMA, "deneme.pdf")


def pdf_olustur(yol, metin):
    akis = f"BT /F1 24 Tf 72 760 Td ({metin}) Tj ET".encode("latin-1")
    nesneler = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(akis) + akis + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    cikti = bytearray(b"%PDF-1.4\n")
    konumlar = []
    for i, nesne in enumerate(nesneler, 1):
        konumlar.append(len(cikti))
        cikti += b"%d 0 obj\n" % i + nesne + b"\nendobj\n"
    xref = len(cikti)
    cikti += b"xref\n0 %d\n0000000000 65535 f \n" % (len(nesneler) + 1)
    for k in konumlar:
        cikti += b"%010d 00000 n \n" % k
    cikti += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(nesneler) + 1, xref)
    with open(yol, "wb") as f:
        f.write(cikti)


if yuklenen.get("openpyxl"):
    @test(K2, "openpyxl: formüllü Excel yazma")
    def _xlsx():
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 21
        ws["A2"] = "=A1*2"
        wb.save(xlsx)
        geri = openpyxl.load_workbook(xlsx)["Sheet"]["A2"].value
        assert geri == "=A1*2", f"formül geri okunamadı: {geri!r}"
        return "dosya yazıldı ve geri okundu"

pdf_olustur(pdf, METIN)

if yuklenen.get("pypdf"):
    @test(K2, "pypdf: PDF metni okuma")
    def _pypdf():
        from pypdf import PdfReader
        metin = PdfReader(pdf).pages[0].extract_text() or ""
        assert "12345" in metin, f"metin bulunamadı: {metin!r}"
        return f"okundu: {metin.strip()!r}"

if yuklenen.get("pdfplumber"):
    @test(K2, "pdfplumber: PDF metni çıkarma")
    def _plumber():
        import pdfplumber
        with pdfplumber.open(pdf) as belge:
            metin = belge.pages[0].extract_text() or ""
        assert "12345" in metin, f"metin bulunamadı: {metin!r}"
        return f"okundu: {metin.strip()!r}"

if yuklenen.get("pypdfium2") and yuklenen.get("PIL.Image"):
    @test(K2, "pypdfium2 + Pillow: PDF sayfasını görüntüye çevirme")
    def _pdfium():
        import pypdfium2 as pdfium
        belge = pdfium.PdfDocument(pdf)
        try:
            goruntu = belge[0].render(scale=1).to_pil()
            png = os.path.join(CALISMA, "sayfa.png")
            goruntu.save(png)
            boyut = goruntu.size
        finally:
            belge.close()
        assert boyut[0] > 100 and os.path.getsize(png) > 0
        return f"sayfa {boyut[0]}x{boyut[1]} piksel görüntüye çevrildi"

if yuklenen.get("cryptography.hazmat.primitives.hashes"):
    @test(K2, "cryptography: derlenmiş bileşen")
    def _crypto():
        from cryptography.hazmat.primitives import hashes
        h = hashes.Hash(hashes.SHA256())
        h.update(b"ksk")
        return f"SHA-256 hesaplandı ({h.finalize().hex()[:12]}…)"


# ------------------------------------------------------------------ Excel otomasyonu
def excel_testi():
    import pythoncom
    import win32api
    import win32con
    import win32event
    import win32process
    import win32com.client

    if not os.path.exists(xlsx):
        bildir(K3, "Excel otomasyonu", "ENGEL", "test dosyası yok (openpyxl adımı başarısız)")
        return
    pythoncom.CoInitialize()
    xl = None
    pid = None
    baslangic = time.time()
    try:
        try:
            xl = win32com.client.DispatchEx("Excel.Application")
        except Exception as exc:  # noqa: BLE001
            bildir(K3, "Excel'i görünmez açma", "ENGEL", hata_ozeti(exc))
            return
        try:
            pid = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
        except Exception:  # noqa: BLE001
            pid = None
        xl.Visible = False
        xl.DisplayAlerts = False
        xl.ScreenUpdating = False
        xl_surum = xl.Version

        wb = xl.Workbooks.Open(xlsx, 0, True)
        xl.CalculateFull()
        deger = wb.Worksheets(1).Range("A2").Value
        wb.Close(False)
        del wb
        if deger == 42:
            bildir(K3, "Görünmez açma + yeniden hesaplama", "GEÇTİ",
                   f"Excel {xl_surum} Python'dan görünmez açıldı, formül hesaplandı (=A1*2 → {deger:g})")
        else:
            bildir(K3, "Görünmez açma + yeniden hesaplama", "UYARI",
                   f"Excel açıldı ama hesap beklenen sonucu vermedi (A2={deger!r})")

        # Mailden/internetten gelmis dosya: Windows'un "internetten indirildi" isareti
        with open(xlsx + ":Zone.Identifier", "w", encoding="ascii") as f:
            f.write("[ZoneTransfer]\nZoneId=3\n")
        try:
            wb2 = xl.Workbooks.Open(xlsx, 0, True)
            d2 = wb2.Worksheets(1).Range("A2").Value
            pv = xl.ProtectedViewWindows.Count
            wb2.Close(False)
            del wb2
            bildir(K3, "İnternetten gelmiş dosyayı açma", "GEÇTİ",
                   f"otomasyonla açıldı (A2={d2!r}, korumalı görünüm penceresi: {pv})")
        except Exception as exc:  # noqa: BLE001
            bildir(K3, "İnternetten gelmiş dosyayı açma", "UYARI", hata_ozeti(exc))
    except Exception as exc:  # noqa: BLE001
        bildir(K3, "Excel otomasyonu", "ENGEL", hata_ozeti(exc))
    finally:
        if xl is not None:
            try:
                xl.Quit()
            except Exception:  # noqa: BLE001
                pass
            del xl
        pythoncom.CoUninitialize()

    # Yalniz bu testin actigi Excel'i bekle; kullanicinin acik Excel'ine dokunma
    if pid:
        try:
            tutamak = win32api.OpenProcess(win32con.SYNCHRONIZE | win32con.PROCESS_TERMINATE, False, pid)
        except Exception:  # noqa: BLE001
            tutamak = None  # surec zaten kapanmis
        if tutamak:
            kapandi = win32event.WaitForSingleObject(tutamak, 20000) == win32event.WAIT_OBJECT_0
            if not kapandi:
                win32api.TerminateProcess(tutamak, 1)
                bildir(K3, "Temiz kapanış", "UYARI",
                       "Excel 20 sn içinde kapanmadı; testin açtığı süreç sonlandırıldı")
            win32api.CloseHandle(tutamak)
            if kapandi:
                bildir(K3, "Temiz kapanış", "GEÇTİ",
                       f"Excel kapandı, arkada süreç kalmadı; toplam {time.time() - baslangic:.1f} sn")
        else:
            bildir(K3, "Temiz kapanış", "GEÇTİ",
                   f"Excel kapandı, arkada süreç kalmadı; toplam {time.time() - baslangic:.1f} sn")


if yuklenen.get("pythoncom") and yuklenen.get("win32com.client"):
    try:
        excel_testi()
    except Exception as exc:  # noqa: BLE001
        bildir(K3, "Excel otomasyonu", "ENGEL", hata_ozeti(exc))
else:
    bildir(K3, "Excel otomasyonu", "ENGEL", "pywin32 yüklenemediği için Excel denenemedi")
